#!/usr/bin/env python3
"""Render the finance report to Word and PDF, preserving correct citations.

Word uses native footnotes (repeated citations can receive new note numbers).
PDF uses the stable unique source index printed in its Sources section.
Requires pandoc, python-docx, and the existing Chrome PDF renderer.
"""
import re
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
src=ROOT/'docs/MONID-FINANCE-TOOLS-REPORT.md'
docx=ROOT/'docs/Monid-Finance-Tools-Report.docx'
subprocess.run(['pandoc',str(src),'-f','gfm','-t','docx','-o',str(docx)],check=True)
doc=Document(docx)
for section in doc.sections:
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.65)
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
for name,size in [('Normal',10),('Heading 1',19),('Heading 2',13),('Heading 3',11),('Heading 4',10),('Footnote Text',8)]:
    if name in doc.styles:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string('202020')
        style.paragraph_format.space_after=Pt(5)
        if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
for table in doc.tables:
    table.autofit=True
    for ri,row in enumerate(table.rows):
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            if ri==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(3)
                for run in paragraph.runs:
                    run.font.size=Pt(8)
                    if ri==0:run.bold=True
for section in doc.sections:
    p=section.footer.paragraphs[0];p.alignment=2;p.add_run('Page ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.save(docx)

# Pandoc emits a fresh note number for some repeated GFM footnotes. Hiding its
# automatic notes while retaining those numbers would break the PDF index.
text=src.read_text()
text=re.sub(r'^\[\^s\d+\]:.*\n?', '',text,flags=re.M)
text=re.sub(r'\[\^s(\d+)\]',r'<sup>\1</sup>',text)
assert '[^s' not in text
with tempfile.TemporaryDirectory(prefix='monid-finance-export-') as tmp:
    pdfsrc=Path(tmp)/'report.md';pdfsrc.write_text(text)
    subprocess.run(['bash',str(ROOT/'scripts/build-report-pdf.sh'),str(pdfsrc),
                    'docs/Monid-Finance-Tools-Report.pdf','--source-list-notes'],cwd=ROOT,check=True)
print('Exported Word with native footnotes and PDF with the stable Sources index.')
